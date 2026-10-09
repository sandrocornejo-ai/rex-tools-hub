"""
migracion_lre_dt_etapa2.py
Migración desde LRE DT — Etapa 2: Creación de archivo de carga de liquidaciones.

Toma el maestro armado en la Etapa 1 (migracion_lre_dt.py) y genera el archivo de importación
"liquidación detalle" de Rex+: una fila por trabajador – mes – concepto.

Reglas: ver claude/especificacion_etapa2_migracion_lre_dt.md (proyecto) — resumen:
  - Columnas del LRE → concepto Rex+ según data/equiv_conceptos.xlsx (montos del mismo concepto se suman).
    4154 → trabajoPesa, 4152 → mutual, 4155 → sis. 4151 → DOS filas (cesAporteCi / cesAporteSol del maestro).
    1116 no genera concepto: se agrega la fila licenciaDias.
  - Conceptos calculados: aporteAFPemp, aporteFAPPCEV, aporteFAPPBAC (no si 1146 = 1 o 1109 = 1) y
    cajaComp (solo FONASA 1143 = 102 con CCAF 1110 ≠ 0; no si 1146 = 1). Base: impo_validado_entrada.
  - Filas con monto 0: se omiten salvo impuesto y cesEmpleado. Licencia mes completo (1115 = 0): solo la
    lista LICENCIA_MES_COMPLETO, aunque vengan en 0.
  - Afecto, Id de institución, Cotización de jubilación, Parcial 7 y Parcial 8 por concepto (tablas abajo).
"""

import io
import os
import re
from datetime import datetime

import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

import migracion_lre_dt as m1

# ─────────────────────────────────────────────
# CONSTANTES
# ─────────────────────────────────────────────
COLUMNAS_CARGA = [
    "Fecha de proceso", "Id empleado", "Número de contrato", "Id del concepto", "Monto del concepto",
    "Afecto", "Id de institución", "Cotización de jubilación", "Días de licencias", "Días trabajados",
    "Fecha de aplicación", "Empresa", "Total de rebajas por LLSS", "Rentas no gravadas",
    "Rebaja por zona extrema", "Jornada", "Días de vacaciones", "Monto Init", "Fase", "Parcial 7", "Parcial 8",
]

COD_4151 = "4151"            # AFC aporte empleador → dos filas (cesAporteCi / cesAporteSol)
COD_DIAS_LIC = "1116"        # Dato: no genera concepto; se agrega la fila licenciaDias
DESTINO_CODIGO = {"4154": "trabajoPesa", "4152": "mutual", "4155": "sis"}   # códigos con varios conceptos

CONCEPTOS_SIEMPRE = {"impuesto", "cesEmpleado"}
LICENCIA_MES_COMPLETO = ["sueldoBase", "gratificacion", "afp", "isapre", "cesEmpleado", "impuesto",
                         "totalesEmpl", "mutual", "sis", "cesAporteSol", "cesAporteCi"]
CALCULADOS_EMPLEADOR = {   # concepto → columna del maestro con la tasa del mes (en %)
    "aporteAFPemp": "aporteAfp_entrada",
    "aporteFAPPCEV": "aporteExpvida_entrada",
    "aporteFAPPBAC": "aporteFappbac_entrada",
}
TASA_CAJA_COMP = "porcAportecaja_entrada"

AFECTO_IMPO_VALIDADO = {"afp", "aporteAFPemp", "aporteFAPPBAC", "cajaComp", "isapre", "mutual",
                        "trabajoPesa", "trabajoPesaEmpl"}
AFECTO_SIS = {"sis", "aporteFAPPCEV"}
AFECTO_BASE_CES_CALC = {"cesAporteCi", "cesAporteSol"}

INSTITUCION_AFP = {"afp", "aporteAFPemp", "sis", "afpAhor", "trabajoPesa", "trabajoPesaEmpl",
                   "cesAporteCi", "cesAporteSol"}
INSTITUCION_CAJA = {"cajaComp", "cajaCred", "cajaAhor", "cajaSegu"}
INSTITUCION_BLANCO = {"aporteFAPPBAC", "cesEmpleado", "licenciaDias"}

PARCIAL7 = {"isapre", "cesAporteSol", "mutual", "aporteAFPemp", "aporteFAPPCEV", "sis"}
PARCIAL8 = {"isapre", "cesAporteSol"}

COL_TECNICO_EXT = "(1146)"
COL_PENSIONADO = "(1109)"
COL_SALUD = "(1143)"
COL_CCAF = "(1110)"
COL_DIAS_TRAB = "(1115)"
COL_DIAS_VAC = "(1117)"
COL_4155 = "(4155)"
COL_5230, COL_5240, COL_3167 = "(5230)", "(5240)", "(3167)"
CODIGO_FONASA = 102
JORNADA_PARCIAL_HORAS = 31

AMARILLO = "FEF3C7"
ETAPA_LOG = "Etapa 2 — Archivo de carga"
CLAVE_GENERAR = "etapa2_generar"


# ─────────────────────────────────────────────
# AUXILIARES
# ─────────────────────────────────────────────
def _num(valor, defecto=0.0):
    """Número o 'defecto' (vacío, texto como 'imponible no encontrado', NaN)."""
    v = pd.to_numeric(valor, errors="coerce")
    return defecto if pd.isna(v) else float(v)


def _entero(valor):
    """Entero redondeado o None si no es numérico."""
    v = pd.to_numeric(valor, errors="coerce")
    return None if pd.isna(v) else int(round(float(v)))


def _texto(valor):
    return "" if valor is None or (not isinstance(valor, str) and pd.isna(valor)) else str(valor).strip()


def _codigo(col):
    """'Sueldo(2101)' / 'Sueldo empresarial (2161)' → '2101' / '2161'. None si la columna no es del LRE."""
    m = re.search(r"\((\d{4})\)\s*$", str(col))
    return m.group(1) if m else None


def cargar_equiv_detalle(ruta=m1.RUTA_EQUIV):
    """
    data/equiv_conceptos.xlsx → {código LRE: [conceptos en el orden del archivo, sin repetir]}.
    Acepta formato nuevo (cod_lre_dt, id_concepto) y antiguo (cod_lre, concepto_detalle).
    Las filas 'Sin codigo en LRE DT' no tienen código y se ignoran.
    """
    if not os.path.exists(ruta):
        raise m1.ErrorArchivo(f"No se encontró el archivo de equivalencias: {ruta}")
    df = pd.read_excel(ruta, dtype=str)
    df.columns = [str(c).strip() for c in df.columns]
    df = df.rename(columns={"cod_lre_dt": "cod_lre", "id_concepto": "concepto_detalle"})
    if "cod_lre" not in df.columns or "concepto_detalle" not in df.columns:
        raise m1.ErrorArchivo("equiv_conceptos.xlsx: faltan las columnas del código LRE o del concepto.")
    equiv = {}
    for cod_txt, concepto in zip(df["cod_lre"], df["concepto_detalle"]):
        cod = _codigo(cod_txt)
        concepto = _texto(concepto)
        if cod and concepto and concepto not in equiv.setdefault(cod, []):
            equiv[cod].append(concepto)
    return equiv


def mapa_codigo_concepto(equiv):
    """Código LRE → un concepto destino (4154/4152/4155 por regla; el resto, el primero del archivo)."""
    mapa = {}
    for cod, conceptos in equiv.items():
        if cod in (COD_4151, COD_DIAS_LIC):
            continue
        mapa[cod] = DESTINO_CODIGO.get(cod, conceptos[0])
    return mapa


def rebajas_por_lista(con_codigo):
    """
    Aportes calculados que la Lista de conceptos trae con Código LRE (validación de la Etapa 1):
    {concepto: código de 4 dígitos}. Solo aplica a aporteAFPemp, aporteFAPPCEV y aporteFAPPBAC.
    """
    rebajas = {}
    if con_codigo is None or con_codigo.empty:
        return rebajas
    for concepto, cod_txt in zip(con_codigo[m1.COL_LC_CONCEPTO], con_codigo[m1.COL_LC_CODIGO_LRE]):
        if concepto in CALCULADOS_EMPLEADOR:
            m = re.search(r"(\d{4})", _texto(cod_txt))
            if m:
                rebajas[concepto] = m.group(1)
    return rebajas


def meses_sin_tope(df):
    """Meses del maestro con tope imponible AFP vacío o 0 → el programa se detiene."""
    tope = pd.to_numeric(df["topeimp_entrada"], errors="coerce").fillna(0)
    return sorted(df.loc[tope <= 0, m1.COL_MES_PROCESO].astype(str).unique().tolist())


# ─────────────────────────────────────────────
# GENERACIÓN DE FILAS
# ─────────────────────────────────────────────
def generar_carga(df, equiv, con_codigo=None):
    """
    df: maestro de la Etapa 1 (solo el rango elegido). equiv: cargar_equiv_detalle().
    Retorna (df_carga, problemas, filas_marcadas):
      problemas      → lista de dicts para el log (Mes, Rut, Fecha inicio, Problema, Valor)
      filas_marcadas → posiciones (0..n-1) de df_carga a pintar en amarillo
    """
    mapa = mapa_codigo_concepto(equiv)
    rebajas = rebajas_por_lista(con_codigo)
    col_rut = m1._buscar_col_rut(df)
    col_ini = m1._buscar_col_codigo(df, m1.COL_FECHA_INI_DT)
    cols_lre = {}                                 # código → nombre real de la columna
    for c in df.columns:
        cod = _codigo(c)
        if cod and cod not in cols_lre:
            cols_lre[cod] = c

    def col(cod_parentesis):
        return m1._buscar_col_codigo(df, cod_parentesis)

    c_lic, c_trab, c_vac = col(m1.COL_DIAS_LIC_DT), col(COL_DIAS_TRAB), col(COL_DIAS_VAC)
    c_5210, c_5220 = col(m1.COL_IMP_TRIB_DT), col(m1.COL_IMP_NO_TRIB_DT)
    c_5230, c_5240, c_3167 = col(COL_5230), col(COL_5240), col(COL_3167)
    c_tec, c_pen, c_salud, c_ccaf, c_4155 = (col(COL_TECNICO_EXT), col(COL_PENSIONADO), col(COL_SALUD),
                                            col(COL_CCAF), col(COL_4155))

    filas, problemas, marcadas = [], [], []

    for _, r in df.iterrows():
        mes = _texto(r[m1.COL_MES_PROCESO])
        rut = _texto(r[col_rut])
        fecha_ini = _texto(r[col_ini]) if col_ini else ""

        def problema(texto, valor=""):
            problemas.append({m1.COL_MES_PROCESO: mes, "Rut trabajador": rut, "Fecha inicio contrato": fecha_ini,
                              "Problema": texto, "Valor buscado": valor})

        g = lambda c: r[c] if c else None   # noqa: E731
        dias_lic, dias_trab, dias_vac = _num(g(c_lic)), _num(g(c_trab)), _num(g(c_vac))
        con_licencia = dias_lic > 0
        lic_mes_completo = dias_trab == 0
        imponible = _num(g(c_5210)) + _num(g(c_5220))
        impo_valid = _entero(r.get(m1.COL_IMPO_VALIDADO))
        empresa = _texto(r.get(m1.COL_ID_EMPRESA))
        horas = pd.to_numeric(r.get(m1.COL_HORAS_SEM), errors="coerce")
        jornada = "" if pd.isna(horas) else ("P" if horas < JORNADA_PARCIAL_HORAS else "C")
        fila_con_problema = False
        if not empresa:
            problema("idEmpresa_entrada vacío: la columna Empresa del archivo de carga queda vacía")
            fila_con_problema = True
        if not jornada:
            problema("horasSema_entrada vacío: no se pudo determinar la Jornada (P/C)")
            fila_con_problema = True
        if impo_valid is None:
            problema("impo_validado_entrada vacío (falta el imponible del mes anterior sin licencia): "
                     "Afecto y aportes calculados quedan en 0")

        # 1) Conceptos calculados (antes de las rebajas, porque el monto se resta de su código LRE)
        calculados = {}
        excluye_empleador = _num(g(c_tec)) == 1 or _num(g(c_pen)) == 1
        if not excluye_empleador:
            for concepto, col_tasa in CALCULADOS_EMPLEADOR.items():
                tasa = _num(r.get(col_tasa))
                calculados[concepto] = int(round(tasa / 100 * (impo_valid or 0)))
        if _num(g(c_salud)) == CODIGO_FONASA and _num(g(c_ccaf)) != 0 and _num(g(c_tec)) != 1:
            calculados["cajaComp"] = int(round(_num(r.get(TASA_CAJA_COMP)) / 100 * (impo_valid or 0)))

        # 2) Rebajas por la Lista de conceptos (aporte calculado que tiene Código LRE asignado)
        ajuste = {}                                  # código LRE → monto a restar
        for concepto, cod in rebajas.items():
            monto_calc = calculados.get(concepto, 0)
            if not monto_calc:
                continue
            if cod.startswith("41"):
                ajuste[cod] = ajuste.get(cod, 0) + monto_calc
            elif cod.startswith("31"):
                problema(f"{concepto} tiene Código LRE {cod} (descuento): no se rebaja nada, revisar",
                         f"Monto calculado: {monto_calc}")

        # 3) Columnas del LRE → conceptos (se suman los del mismo concepto)
        montos = {}
        for cod, c in cols_lre.items():
            valor = _num(r[c])
            if cod in ajuste:
                valor -= ajuste[cod]
                if valor < 0:
                    problema(f"Columna {cod} queda negativa después de rebajar aportes calculados", f"{valor:.0f}")
            if cod in (COD_4151, COD_DIAS_LIC):
                continue
            if cod in mapa:
                montos[mapa[cod]] = montos.get(mapa[cod], 0) + valor
            elif cod[0] in "234" and valor != 0:
                problema(f"Columna {c} tiene monto y no está en equiv_conceptos.xlsx: no genera fila", f"{valor:.0f}")

        # 4) AFC aporte empleador (4151) desglosado en la Etapa 1
        ci, sol = _entero(r.get(m1.COL_CES_APORTE_CI)), _entero(r.get(m1.COL_CES_APORTE_SOL))
        if ci is None or sol is None:
            if _num(g(cols_lre.get(COD_4151))) != 0:
                problema("4151 sin desglose (factor_ces_entrada vacío): no se generan cesAporteCi / cesAporteSol",
                         f"4151: {_num(g(cols_lre.get(COD_4151))):.0f}")
        else:
            montos["cesAporteCi"] = montos.get("cesAporteCi", 0) + ci
            montos["cesAporteSol"] = montos.get("cesAporteSol", 0) + sol

        for concepto, monto in calculados.items():
            montos[concepto] = montos.get(concepto, 0) + monto

        # 5) Conceptos que van siempre
        forzados = set(CONCEPTOS_SIEMPRE) | (set(LICENCIA_MES_COMPLETO) if lic_mes_completo else set())
        for concepto in sorted(forzados):
            montos.setdefault(concepto, 0)

        # 6) Valores comunes del registro
        contrato = r.get(m1.COL_NUM_CONTRATO)
        base = {
            "Fecha de proceso": mes, "Id empleado": rut,
            "Número de contrato": None if contrato is None or pd.isna(contrato) else contrato,
            "Días de licencias": _entero(dias_lic) or 0, "Días trabajados": _entero(dias_trab) or 0,
            "Fecha de aplicación": mes, "Empresa": empresa, "Jornada": jornada,
            "Días de vacaciones": _entero(dias_vac) or 0, "Fase": 1,
        }
        afp, salud, caja = (_texto(r.get(m1.COL_NOMBRE_AFP)), _texto(r.get(m1.COL_NOMBRE_SALUD)),
                            _texto(r.get(m1.COL_NOMBRE_CAJA)))
        ult_imp = _entero(r.get(m1.COL_ULT_IMP_SIN_LIC))
        baseces_valid = _entero(r.get(m1.COL_BASECES_VALIDADO))
        base_ces_calc = _entero(r.get(m1.COL_BASE_CES_CALC))

        for concepto, monto in montos.items():
            if lic_mes_completo and concepto not in LICENCIA_MES_COMPLETO:
                continue
            if monto == 0 and concepto not in forzados:
                continue
            monto = int(round(monto))

            # Afecto
            if concepto in AFECTO_IMPO_VALIDADO:
                afecto = impo_valid or 0
            elif concepto in AFECTO_SIS:
                if con_licencia:
                    sis = _num(r.get(m1.COL_SIS))
                    if sis > 0:
                        afecto = int(round(_num(g(c_4155)) * 100 / sis))
                    else:
                        afecto = 0
                        problema(f"Afecto de {concepto}: Sis_entrada vacío o 0 con licencia → Afecto 0")
                else:
                    afecto = impo_valid or 0
            elif concepto in AFECTO_BASE_CES_CALC:
                afecto = base_ces_calc or 0
            elif concepto == "cesEmpleado":
                afecto = baseces_valid or 0
            elif concepto == "totalesEmpl":
                afecto = int(round(imponible))
            elif concepto == "impuesto":
                afecto = int(round(imponible - _num(r.get(m1.COL_REBAJA_LLSS))))
            else:
                afecto = 0

            # Id de institución
            if concepto in INSTITUCION_AFP:
                institucion = afp
            elif concepto == "isapre":
                institucion = salud
            elif concepto == "impuesto":
                institucion = "Impuesto"
            elif concepto == "mutual":
                institucion = _texto(r.get(m1.COL_NOMBRE_MUTUAL))
            elif concepto == "aporteFAPPCEV":
                institucion = "seguridadsocial"
            elif concepto in INSTITUCION_CAJA:
                institucion = caja
            elif concepto == "apvi":
                institucion = f"apv{afp}" if afp else ""
            elif concepto in INSTITUCION_BLANCO:
                institucion = ""
            else:
                institucion = 0

            # Cotización de jubilación
            cotizacion = {
                "afp": r.get(m1.COL_PORC_AFP),
                "aporteAFPemp": r.get("aporteAfp_entrada"),
                "aporteFAPPBAC": r.get("aporteFappbac_entrada"),
                "aporteFAPPCEV": r.get("aporteExpvida_entrada"),
                "apvi": monto, "isapre": monto, "impuesto": "",
                "mutual": r.get(m1.COL_PORC_MUTUAL), "sis": r.get(m1.COL_SIS),
            }.get(concepto, 0)
            if concepto == "totalesEmpl":   # min(5210 + 5220, topeimp_entrada)
                tope = _num(r.get("topeimp_entrada"))
                cotizacion = int(round(min(imponible, tope))) if tope > 0 else int(round(imponible))

            # Parciales
            if concepto in PARCIAL7 and con_licencia:
                if ult_imp is None:
                    parcial7 = ""
                    problema(f"Parcial 7 de {concepto}: ultImpSinLic_entrada vacío o 'imponible no encontrado'")
                else:
                    parcial7 = ult_imp
            else:
                parcial7 = 0
            parcial8 = (baseces_valid or 0) if concepto in PARCIAL8 else 0

            es_impuesto = concepto == "impuesto"
            filas.append({
                **base,
                "Id del concepto": concepto, "Monto del concepto": monto, "Afecto": afecto,
                "Id de institución": institucion, "Cotización de jubilación": cotizacion,
                "Total de rebajas por LLSS": (_entero(r.get(m1.COL_REBAJA_LLSS)) or 0) if es_impuesto else 0,
                "Rentas no gravadas": int(round(_num(g(c_5230)) + _num(g(c_5240)))) if es_impuesto else 0,
                "Rebaja por zona extrema": int(round(_num(g(c_3167)))) if es_impuesto else 0,
                "Monto Init": (_entero(r.get(m1.COL_SUELDO_CONTRATO)) or 0) if concepto == "sueldoBase" else 0,
                "Parcial 7": parcial7, "Parcial 8": parcial8,
            })
            if fila_con_problema:
                marcadas.append(len(filas) - 1)

        # 7) Fila adicional licenciaDias
        if con_licencia:
            filas.append({
                **base,
                "Id del concepto": "licenciaDias", "Monto del concepto": _entero(dias_lic), "Afecto": 0,
                "Id de institución": "", "Cotización de jubilación": "",
                "Total de rebajas por LLSS": 0, "Rentas no gravadas": 0, "Rebaja por zona extrema": 0,
                "Monto Init": 0, "Parcial 7": 0, "Parcial 8": 0,
            })
            if fila_con_problema:
                marcadas.append(len(filas) - 1)

    df_carga = pd.DataFrame(filas, columns=COLUMNAS_CARGA)
    return df_carga, problemas, marcadas


# ─────────────────────────────────────────────
# EXCEL DE SALIDA
# ─────────────────────────────────────────────
def excel_carga(df_carga, marcadas=()):
    """Excel con formato Rex+ (igual que modulo_dt). Filas con problemas en amarillo."""
    output = io.BytesIO()
    wb = Workbook()
    ws = wb.active
    ws.title = "Liquidaciones"
    relleno_enc = PatternFill("solid", fgColor="1A2744")
    fuente_enc = Font(bold=True, color="FFFFFF", size=10)
    borde = Border(bottom=Side(style="thin", color="E8EDF5"), right=Side(style="thin", color="E8EDF5"))
    par, impar = PatternFill("solid", fgColor="EAF0F8"), PatternFill("solid", fgColor="FFFFFF")
    amarillo = PatternFill("solid", fgColor=AMARILLO)
    marcadas = set(marcadas)
    for ci, nombre in enumerate(df_carga.columns, 1):
        celda = ws.cell(row=1, column=ci, value=nombre)
        celda.fill, celda.font = relleno_enc, fuente_enc
        celda.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[celda.column_letter].width = max(len(nombre) + 4, 14)
    for pos, fila in enumerate(df_carga.itertuples(index=False)):
        ri = pos + 2
        relleno = amarillo if pos in marcadas else (par if ri % 2 == 0 else impar)
        for ci, valor in enumerate(fila, 1):
            if valor is not None and not isinstance(valor, str) and pd.isna(valor):
                valor = None
            celda = ws.cell(row=ri, column=ci, value=valor)
            celda.fill, celda.border = relleno, borde
            celda.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"
    wb.save(output)
    return output.getvalue()


# ─────────────────────────────────────────────
# PANTALLA (se llama desde migracion_lre_dt.main, después del maestro)
# ─────────────────────────────────────────────
def seccion_etapa2(df_cons, con_codigo, log, mes_desde, mes_hasta):
    st.divider()
    st.subheader("📤 Etapa 2 — Creación de archivo de carga de liquidaciones")
    st.caption("Genera el archivo de importación 'liquidación detalle' de Rex+ a partir del maestro de arriba "
               "(incluye los datos ingresados a mano).")
    if st.button("⚙️ Generar archivo de carga de liquidaciones", type="primary", key="btn_etapa2"):
        st.session_state[CLAVE_GENERAR] = True
    if not st.session_state.get(CLAVE_GENERAR):
        return

    sin_tope = meses_sin_tope(df_cons)
    if sin_tope:
        msg = ("El tope imponible AFP (topeimp_entrada) está vacío o en 0 para: **" + ", ".join(sin_tope)
               + "**. Completa esos meses en Parámetros mensuales y vuelve a subir el archivo.")
        log.error(ETAPA_LOG, msg.replace("**", ""))
        st.error(f"❌ {msg}")
        return
    try:
        equiv = cargar_equiv_detalle()
        df_carga, problemas, marcadas = generar_carga(df_cons, equiv, con_codigo)
    except Exception as e:  # noqa: BLE001
        log.error(ETAPA_LOG, f"No se pudo generar el archivo de carga: {e}")
        st.error(f"❌ No se pudo generar el archivo de carga: {e}")
        return

    # Log: un resumen por tipo de problema + detalle por registro
    if problemas:
        log.detalle.extend(problemas)
        por_tipo = pd.Series([re.sub(r"\(\d{4}\)|\d{4}", "####", p["Problema"]) for p in problemas]).value_counts()
        for tipo, n in por_tipo.items():
            log.advertencia(ETAPA_LOG, f"{n} caso(s): {tipo}")
        st.warning(f"⚠️ {len(problemas)} observación(es) en la generación (filas afectadas en amarillo en el Excel). "
                   "Revisa la hoja 'Detalle registros' del log.")
        with st.expander("Ver observaciones"):
            st.dataframe(pd.DataFrame(problemas), use_container_width=True, hide_index=True)
    else:
        log.info(ETAPA_LOG, "Archivo de carga generado sin observaciones.")

    n_trab = df_carga["Id empleado"].nunique()
    st.success(f"✅ **Archivo de carga generado** — {len(df_carga)} filas · {n_trab} trabajador(es) · "
               f"meses {mes_desde} a {mes_hasta}" + ("" if not problemas else " (con observaciones)"))
    log.info(ETAPA_LOG, f"{len(df_carga)} filas, {n_trab} trabajadores, meses {mes_desde} a {mes_hasta}.")
    resumen = (df_carga.groupby("Id del concepto")
               .agg(Filas=("Monto del concepto", "size"), Monto=("Monto del concepto", "sum"))
               .reset_index().sort_values("Id del concepto"))
    with st.expander("📊 Resumen por concepto"):
        st.dataframe(resumen, use_container_width=True, hide_index=True)
    with st.expander("👁️ Vista previa del archivo de carga"):
        st.dataframe(df_carga.head(200), use_container_width=True, hide_index=True)
    st.download_button(
        "⬇️ Descargar archivo de carga de liquidaciones (.xlsx)",
        data=excel_carga(df_carga, marcadas),
        file_name=f"carga_liquidaciones_{mes_desde}_a_{mes_hasta}_{datetime.now():%Y%m%d_%H%M%S}.xlsx",
        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        key="dl_etapa2",
    )
